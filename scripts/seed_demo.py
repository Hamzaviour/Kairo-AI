"""
Demo Database Bootstrap & Seed Script
=====================================
- Applies backend/migrations/001_initial_schema.sql when the schema is missing
- Seeds the two demo accounts and their tasks from data/Tasks_and_Users_Updated.xlsx
- Adds a handful of extra demo tasks (with subtasks) so every view has content
- Idempotent: re-running skips rows that already exist

Usage (from the repository root):
    DATABASE_URL=postgresql://... python scripts/seed_demo.py
"""

import os
import sys
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from sqlalchemy import inspect, text

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "backend"))

from app.core.database import engine, SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.models import *  # noqa: E402,F401,F403  (registers every mapper)
from app.models.user import User, UserPreferences  # noqa: E402
from app.models.task import Task, Subtask  # noqa: E402
from app.models.project import Project, ProjectMember  # noqa: E402
from app.models.collaboration import TaskActivityLog  # noqa: E402

SCHEMA_FILE = REPO_ROOT / "backend" / "migrations" / "001_initial_schema.sql"
EXCEL_FILE = REPO_ROOT / "data" / "Tasks_and_Users_Updated.xlsx"

STATUS_MAP = {
    "To Do": "todo",
    "Pending": "todo",
    "In Progress": "in_progress",
    "In Review": "in_review",
    "Under Review": "in_review",
    "Blocked": "blocked",
    "Completed": "completed",
}

# Extra demo tasks: (title, description, status, priority, owner email, due in days, subtasks)
EXTRA_TASKS = [
    ("Prepare Q4 budget report", "Consolidate department spend and submit the Q4 budget report for review.",
     "todo", "urgent", "alexander.wright@enterprise.com", 3,
     ["Collect department spend sheets", "Reconcile vendor invoices", "Draft executive summary", "Submit for approval"]),
    ("Migrate legacy spreadsheet tracker", "Import the remaining Excel task sheets into TaskPilot and retire the old file.",
     "in_review", "high", "alexander.wright@enterprise.com", 5,
     ["Export sheets to .xlsx", "Run import dry-run preview", "Commit validated rows"]),
    ("Write onboarding guide", "Short guide for new teammates covering views, AI Copilot and Excel import.",
     "backlog", "low", "alexander.wright@enterprise.com", 14, []),
    ("Fix flaky login redirect", "Users are occasionally bounced to the login page after token refresh.",
     "blocked", "high", "sarah.jenkins@enterprise.com", 2,
     ["Reproduce with expired access token", "Add refresh-queue regression test"]),
    ("Design Kanban empty states", "Illustrations and copy for empty Kanban columns.",
     "in_progress", "medium", "sarah.jenkins@enterprise.com", 6,
     ["Sketch three variants", "Review with design lead"]),
    ("Weekly sprint retro notes", "Summarise wins, misses and action items from the sprint retro.",
     "todo", "medium", "sarah.jenkins@enterprise.com", 1, []),
]


def ensure_schema() -> None:
    if inspect(engine).has_table("users"):
        print("Schema present - skipping migration")
        return
    print(f"Applying {SCHEMA_FILE.name} ...")
    raw = engine.raw_connection()
    try:
        with raw.cursor() as cur:
            cur.execute(SCHEMA_FILE.read_text(encoding="utf-8"))
        raw.commit()
    finally:
        raw.close()


def load_excel():
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
    users = [r for r in wb["Users"].iter_rows(min_row=2, values_only=True) if r[0]]
    tasks = [r for r in wb["Task"].iter_rows(min_row=2, values_only=True) if r[0]]
    return users, tasks


def seed() -> None:
    ensure_schema()
    users, tasks = load_excel()
    now = datetime.now(timezone.utc)
    db = SessionLocal()
    try:
        by_email = {}
        for uid, name, email, password, pic, role, _ in users:
            user = db.query(User).filter(User.email == email).first()
            if not user:
                user = User(id=uuid.UUID(uid), name=name, email=email, password_hash=hash_password(password),
                            role=role or "user", profile_pic=pic)
                db.add(user)
                db.add(UserPreferences(user_id=user.id))
                print(f"+ user {email}")
            by_email[email] = user
        db.flush()

        for i, (tid, title, desc, url, image, status, priority, owner_id) in enumerate(tasks):
            if db.get(Task, uuid.UUID(tid)):
                continue
            # The workbook's image URLs point at placeholder CDNs, so they are not carried over
            db.add(Task(id=uuid.UUID(tid), title=title, description=desc, url=url,
                        status=STATUS_MAP.get(status, "todo"), priority=(priority or "medium").lower(),
                        creator_id=uuid.UUID(owner_id), assignee_id=uuid.UUID(owner_id),
                        due_date=now + timedelta(days=4 + i * 3), position=i,
                        completed_at=now if STATUS_MAP.get(status) == "completed" else None))
            print(f"+ task {title}")

        for i, (title, desc, status, priority, email, due, subtasks) in enumerate(EXTRA_TASKS):
            if db.query(Task).filter(Task.title == title).first():
                continue
            owner = by_email[email]
            task = Task(title=title, description=desc, status=status, priority=priority,
                        creator_id=owner.id, assignee_id=owner.id, due_date=now + timedelta(days=due),
                        estimated_minutes=60 * (i % 3 + 1), position=10 + i)
            task.subtasks = [Subtask(title=s, position=j, is_completed=j == 0, completed_at=now if j == 0 else None)
                             for j, s in enumerate(subtasks)]
            db.add(task)
            print(f"+ task {title}")

        db.flush()

        # One demo project grouping the admin's migration work. Sarah is deliberately not a member:
        # membership widens task scope, and the demo account must stay strictly isolated.
        project = db.query(Project).filter(Project.name == "Spreadsheet Migration").first()
        if not project:
            admin = by_email["alexander.wright@enterprise.com"]
            project = Project(name="Spreadsheet Migration", color="#6366f1", owner_id=admin.id,
                              description="Move the team's Excel task tracker into TaskPilot and retire the legacy sheet.")
            project.members = [ProjectMember(user_id=admin.id, role="owner")]
            db.add(project)
            db.flush()
            for title in ("Migrate legacy spreadsheet tracker", "Prepare Q4 budget report",
                          "Implement User Authentication", "Setup CI/CD Pipeline"):
                db.query(Task).filter(Task.title == title).update({"project_id": project.id})
            print("+ project Spreadsheet Migration")

        # Creation entries so the audit trail reflects the seeded work
        for task in db.query(Task).filter(~Task.activity_logs.any()).all():
            db.add(TaskActivityLog(task_id=task.id, user_id=task.creator_id, action="created", new_value=task.title))

        db.commit()
        count = db.execute(text("SELECT count(*) FROM tasks")).scalar()
        print(f"Seed complete - {count} tasks in database")
    finally:
        db.close()


if __name__ == "__main__":
    if not os.environ.get("DATABASE_URL"):
        sys.exit("DATABASE_URL is not set")
    seed()
