import io
import openpyxl

def test_excel_export(client, user_token):
    res = client.get("/api/excel/export", headers={"Authorization": f"Bearer {user_token}"})
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert len(res.content) > 1000

    # Parse exported workbook to verify integrity
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Tasks" in wb.sheetnames
    ws = wb["Tasks"]
    assert ws.max_row >= 2  # Header + at least Sarah's task
    headers = [cell.value for cell in ws[1]]
    assert "Task ID" in headers
    assert "Title" in headers
    assert "Status" in headers
    assert "Priority" in headers

def test_excel_import_preview(client, user_token):
    # Create test in-memory workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task"
    ws.append(["Task Heading", "Task Description", "Task Status", "Task Priority"])
    ws.append(["Imported Test Task", "Created via automated test", "In Progress", "High"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    files = {"file": ("test_import.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post("/api/excel/preview-import", files=files, headers={"Authorization": f"Bearer {user_token}"})
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["valid_rows_count"] == 1
    assert data["preview_items"][0]["title"] == "Imported Test Task"
    assert data["preview_items"][0]["status"] == "in_progress"
    assert data["preview_items"][0]["priority"] == "high"
