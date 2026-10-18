import io
from typing import List, Tuple, Optional
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.task import Task
from app.models.user import User
from app.services.task_service import get_scoped_tasks_query
from app.schemas.excel import ExcelTaskPreviewRow, ExcelImportPreviewResponse

class ExcelService:
    @staticmethod
    def export_tasks_to_excel(db: Session, current_user: User) -> io.BytesIO:
        """
        Exports all permitted tasks for the current user into a professionally styled Excel workbook.
        Includes Status and Priority dropdown validations.
        """
        tasks = get_scoped_tasks_query(db, current_user).order_by(Task.position.asc()).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tasks"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Task ID", "Title", "Description", "URL", "Image URL",
            "Status", "Priority", "Assignee Name", "Assignee Email",
            "Due Date", "Created Date"
        ]
        ws.append(headers)

        # Style headers
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Data rows
        data_font = Font(name="Calibri", size=10)
        mono_font = Font(name="Consolas", size=9, color="555555")

        for r_idx, t in enumerate(tasks, start=2):
            assignee_name = t.assignee.name if t.assignee else "Unassigned"
            assignee_email = t.assignee.email if t.assignee else ""
            due_str = t.due_date.strftime("%Y-%m-%d %H:%M") if t.due_date else ""
            created_str = t.created_at.strftime("%Y-%m-%d %H:%M") if t.created_at else ""

            row_data = [
                str(t.id),
                t.title,
                t.description or "",
                t.url or "",
                t.image_url or "",
                t.status,
                t.priority,
                assignee_name,
                assignee_email,
                due_str,
                created_str
            ]
            ws.append(row_data)
            ws.row_dimensions[r_idx].height = 22

            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = mono_font if c_idx == 1 else data_font
                cell.border = thin_border
                if c_idx in [6, 7]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Data validations
        status_dv = DataValidation(
            type="list",
            formula1='"backlog,todo,in_progress,in_review,blocked,completed"',
            allow_blank=True
        )
        ws.add_data_validation(status_dv)
        status_dv.add("F2:F5000")

        priority_dv = DataValidation(
            type="list",
            formula1='"low,medium,high,urgent"',
            allow_blank=True
        )
        ws.add_data_validation(priority_dv)
        priority_dv.add("G2:G5000")

        # Column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

        # Output to memory buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def preview_excel_import(file_contents: bytes, filename: str) -> ExcelImportPreviewResponse:
        """
        Parses an uploaded Excel file, validates headers and rows, and returns a preview report.
        """
        buffer = io.BytesIO(file_contents)
        wb = openpyxl.load_workbook(buffer, data_only=True)

        tasks_sheet_name = None
        for name in wb.sheetnames:
            if "task" in name.lower():
                tasks_sheet_name = name
                break
        if not tasks_sheet_name:
            tasks_sheet_name = wb.sheetnames[0]

        ws = wb[tasks_sheet_name]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        
        # Fuzzy header matching
        col_map = {}
        for idx, h in enumerate(headers):
            hl = h.lower()
            if "title" in hl or "heading" in hl:
                col_map["title"] = idx
            elif "desc" in hl:
                col_map["description"] = idx
            elif "status" in hl:
                col_map["status"] = idx
            elif "priority" in hl:
                col_map["priority"] = idx
            elif "url" in hl and "image" not in hl:
                col_map["url"] = idx
            elif "image" in hl or "pic" in hl:
                col_map["image"] = idx
            elif "user" in hl or "assign" in hl or "email" in hl:
                col_map["user"] = idx

        preview_rows: List[ExcelTaskPreviewRow] = []
        valid_count = 0
        invalid_count = 0

        valid_statuses = {"backlog", "todo", "in_progress", "in_review", "blocked", "completed"}
        valid_priorities = {"low", "medium", "high", "urgent"}

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            
            title_idx = col_map.get("title", 0)
            title = str(row[title_idx] or "").strip() if title_idx < len(row) else ""
            
            desc_idx = col_map.get("description")
            desc = str(row[desc_idx] or "").strip() if desc_idx is not None and desc_idx < len(row) and row[desc_idx] else None

            status_idx = col_map.get("status")
            raw_status = str(row[status_idx] or "todo").strip().lower().replace(" ", "_") if status_idx is not None and status_idx < len(row) and row[status_idx] else "todo"
            if raw_status == "pending":
                raw_status = "todo"
            if raw_status not in valid_statuses:
                raw_status = "todo"

            prio_idx = col_map.get("priority")
            raw_priority = str(row[prio_idx] or "medium").strip().lower() if prio_idx is not None and prio_idx < len(row) and row[prio_idx] else "medium"
            if raw_priority not in valid_priorities:
                raw_priority = "medium"

            url_idx = col_map.get("url")
            url = str(row[url_idx] or "").strip() if url_idx is not None and url_idx < len(row) and row[url_idx] else None

            img_idx = col_map.get("image")
            img = str(row[img_idx] or "").strip() if img_idx is not None and img_idx < len(row) and row[img_idx] else None

            user_idx = col_map.get("user")
            user_ref = str(row[user_idx] or "").strip() if user_idx is not None and user_idx < len(row) and row[user_idx] else None

            is_valid = True
            err = None
            if not title:
                is_valid = False
                err = "Task title cannot be empty"
                invalid_count += 1
            else:
                valid_count += 1

            preview_rows.append(
                ExcelTaskPreviewRow(
                    row_number=row_num,
                    title=title,
                    description=desc,
                    url=url,
                    image_url=img,
                    status=raw_status,
                    priority=raw_priority,
                    user_email_or_name=user_ref,
                    is_valid=is_valid,
                    validation_error=err
                )
            )

        return ExcelImportPreviewResponse(
            filename=filename,
            total_rows_detected=len(preview_rows),
            valid_rows_count=valid_count,
            invalid_rows_count=invalid_count,
            users_sheet_found="Users" in wb.sheetnames,
            tasks_sheet_found=tasks_sheet_name is not None,
            preview_items=preview_rows,
            warnings=[] if invalid_count == 0 else [f"{invalid_count} row(s) contain missing or invalid titles and will be skipped"]
        )
